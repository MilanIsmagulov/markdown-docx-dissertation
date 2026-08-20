from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm
from openpyxl import Workbook

from builder.assets import _read_rows, _table_config_marker, process_assets
from builder.docx_renderer import _format_advanced_tables
from builder.resolver import build_index
from builder.validator import validate_index


def test_csv_range_selects_rows_and_columns(tmp_path: Path) -> None:
    source = tmp_path / "data.csv"
    source.write_text("A,B,C,D\n1,2,3,4\n5,6,7,8\n", encoding="utf-8")

    assert _read_rows(source, None, "B2:C3") == [["2", "3"], ["6", "7"]]


def test_xlsx_sheet_and_range_are_selected(tmp_path: Path) -> None:
    source = tmp_path / "data.xlsx"
    workbook = Workbook()
    workbook.active.title = "Other"
    sheet = workbook.create_sheet("Experiment")
    sheet.append(["A", "B", "C"])
    sheet.append([1, 2, 3])
    workbook.save(source)

    assert _read_rows(source, "Experiment", "B1:C2") == [["B", "C"], ["2", "3"]]


def test_table_directive_emits_continuations_notes_and_landscape_section(tmp_path: Path) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    (assets / "data.csv").write_text(
        "A,B,C,D,E,F,G\n" + "\n".join(f"{n},2,3,4,5,6,7" for n in range(1, 6)),
        encoding="utf-8",
    )
    markdown = """# Глава 1
```table
source: assets/data.csv
id: results
caption: Результаты
split_rows: 2
section: auto
widths: [30mm, 30mm, 30mm, 30mm, 30mm, 30mm, 30mm]
align: [left, center, center, center, center, center, right]
merges: [A1:B1]
note: Демонстрационное примечание
source_note: Составлено автором
```
"""

    result = process_assets(markdown, tmp_path)

    assert result.count("[[TABLE_CONFIG:") == 3
    assert "[[SECTION:landscape]]" in result
    assert "[[SECTION:dissertation]]" in result
    assert "Продолжение таблицы [[NUMBER:table:results:1.1]]" in result
    assert "Примечание. Демонстрационное примечание" in result
    assert "Источник: Составлено автором" in result


def test_docx_table_config_applies_widths_alignment_merges_and_header_flag() -> None:
    document = Document()
    document.sections[0].page_width = Mm(210)
    document.sections[0].left_margin = Mm(25)
    document.sections[0].right_margin = Mm(10)
    document.add_paragraph(_table_config_marker({
        "widths": ["50mm", "60mm", "65mm"],
        "align": ["left", "center", "right"],
        "merges": ["A1:B1"],
        "repeat_header": False,
    }))
    table = document.add_table(rows=2, cols=3)
    for column, text in enumerate(("One", "Two", "Three")):
        table.cell(0, column).text = text
        table.cell(1, column).text = str(column)
    header = table.rows[0]._tr.get_or_add_trPr()
    from docx.oxml import OxmlElement

    header.append(OxmlElement("w:tblHeader"))

    _format_advanced_tables(document)

    assert not any(paragraph.text.startswith("[[TABLE_CONFIG:") for paragraph in document.paragraphs)
    assert table.cell(0, 0)._tc is table.cell(0, 1)._tc
    assert table.rows[1].cells[0].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert table.rows[1].cells[1].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert table.rows[1].cells[2].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert not table.rows[0]._tr.get_or_add_trPr().findall(qn("w:tblHeader"))
    assert [int(node.get(qn("w:w"))) for node in table._tbl.tblGrid] == [2834, 3401, 3685]


def test_table_validator_reports_invalid_options_and_wide_portrait_table(tmp_path: Path) -> None:
    content = tmp_path / "content"
    assets = content / "assets"
    assets.mkdir(parents=True)
    (assets / "data.csv").write_text("A,B\n1,2\n", encoding="utf-8")
    root = content / "root.md"
    root.write_text(
        """---
id: document:root
type: document
title: Root
---
```table
source: assets/data.csv
id: data
caption: Data
sheet: Wrong
range: broken
widths: [100mm, 100mm]
align: [middle]
merges: [bad]
repeat_header: "yes"
split_rows: 0
section: portrait
```
""",
        encoding="utf-8",
    )

    diagnostics = validate_index(build_index(tmp_path, content), root)
    codes = {item.code for item in diagnostics}

    assert {"E_TABLE_SHEET", "E_TABLE_RANGE", "W_TABLE_WIDE", "E_TABLE_ALIGN"} <= codes
    assert {"E_TABLE_MERGES", "E_TABLE_REPEAT_HEADER", "E_TABLE_SPLIT"} <= codes
