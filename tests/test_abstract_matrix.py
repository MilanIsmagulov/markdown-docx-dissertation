import json
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from builder.assets import process_assets
from builder.docx_renderer import _format_publication_paragraphs, _prepend_abstract_front_matter


SNAPSHOT = Path(__file__).parent / "snapshots" / "abstract-front-matter.json"


def _metadata(*, long: bool = False) -> dict:
    suffix = " с дополнительным демонстрационным продолжением" if long else ""
    return {
        "title": "Название диссертационной работы" + suffix,
        "degree": "кандидата технических наук",
        "author": {"full_name": "Иванов Иван Иванович" + suffix},
        "specialty": {"code": "2.3.1", "name": "Системный анализ, управление и обработка информации, статистика"},
        "supervisor": {
            "full_name": "Петров Петр Петрович" + suffix,
            "degree": "д-р техн. наук", "title": "профессор",
        },
        "city": "Город", "year": 2000,
    }


def _config(opponents: list[dict] | None = None, *, partial: bool = False) -> dict:
    defense = {
        "organization": "ФГБОУ ВО «Название университета»",
        "organization_unit": "наименование структурного подразделения",
        "opponents": opponents or [],
        "leading_organization": "[указать ведущую организацию]",
        "date": "«_____» _____________ 20__ г.", "time": "_________",
        "council": "[указать совет]", "address": "[указать адрес]",
        "library": "[указать библиотеку]", "website": "[указать сайт]",
        "mailing_date": "«_____» _____________ 20__ г",
        "secretary": "[указать секретаря]", "secretary_degree": "[указать степень]",
    }
    if partial:
        defense = {"organization": defense["organization"], "opponents": opponents or []}
    return {
        "layout": {"title_font_size": "11pt", "verso_font_size": "10pt"},
        "defense": defense,
    }


def _front_matter(metadata: dict, config: dict) -> Document:
    document = Document()
    document.add_paragraph("BODY")
    _prepend_abstract_front_matter(document, metadata, config)
    return document


def _signature(document: Document) -> dict:
    titles = []
    for paragraph in document.paragraphs[:8]:
        run = paragraph.runs[0]
        titles.append({
            "text": paragraph.text,
            "alignment": paragraph.alignment,
            "after_pt": round(paragraph.paragraph_format.space_after.pt, 2),
            "bold": bool(run.bold),
            "size_pt": round(run.font.size.pt, 2),
        })
    table = document.tables[0]
    rows = []
    for row in table.rows:
        unique = []
        seen = set()
        for cell in row.cells:
            identity = id(cell._tc)
            if identity not in seen:
                unique.append(cell.text)
                seen.add(identity)
        rows.append(unique)
    return {
        "title": titles,
        "verso_rows": rows,
        "verso_row_count": len(rows),
        "table_width": table._tbl.tblPr.xpath("./w:tblW/@w:w")[0],
        "table_layout": table._tbl.tblPr.xpath("./w:tblLayout/@w:type")[0],
        "sections": len(document.sections),
        "body_page_start": document.sections[-1]._sectPr.xpath("./w:pgNumType/@w:start"),
    }


def test_title_and_verso_match_structural_snapshot() -> None:
    opponents = [
        {"degree": "канд. техн. наук, проф.", "full_name": "ИВАНОВ Иван Иванович"},
        {"degree": "д-р техн. наук, проф.", "full_name": "СИДОРОВ Сидор Сидорович"},
    ]
    actual = _signature(_front_matter(_metadata(), _config(opponents)))
    expected = json.loads(SNAPSHOT.read_text(encoding="utf-8"))
    assert actual == expected


@pytest.mark.parametrize("count", [0, 1, 3])
def test_verso_supports_zero_one_and_many_opponents(count: int) -> None:
    opponents = [
        {"degree": f"Учёная степень {index}", "full_name": f"ИВАНОВ Иван Иванович {index}"}
        for index in range(1, count + 1)
    ]
    document = _front_matter(_metadata(), _config(opponents))
    table = document.tables[0]
    text = "\n".join(cell.text for row in table.rows for cell in row.cells)

    assert len(table.rows) == 8 + 2 * count
    assert text.count("Официальные оппоненты:") == 2 * count
    for opponent in opponents:
        assert opponent["full_name"] in text


def test_long_metadata_keeps_autofit_grid_without_fixed_row_heights() -> None:
    long_name = "ИВАНОВ Иван Иванович с очень длинным демонстрационным составным именем"
    document = _front_matter(
        _metadata(long=True),
        _config([{"degree": "д-р технических наук, профессор", "full_name": long_name}]),
    )
    table = document.tables[0]

    assert long_name in "\n".join(cell.text for row in table.rows for cell in row.cells)
    assert table.autofit is True
    assert table._tbl.tblPr.xpath("./w:tblLayout/@w:type") == ["autofit"]
    assert not table._tbl.xpath(".//w:trHeight")


def test_partially_filled_defense_uses_safe_placeholders() -> None:
    document = _front_matter(_metadata(), _config(partial=True))
    text = "\n".join(cell.text for row in document.tables[0].rows for cell in row.cells)

    assert "[указать дату]" in text
    assert "[указать время]" in text
    assert "[указать секретаря]" in text


def test_large_publication_list_keeps_numbering_and_hanging_indent(tmp_path: Path) -> None:
    bibliography = tmp_path / "publications.bib"
    bibliography.write_text(
        "\n".join(
            f"@article{{item{number}, author={{Иванов, Иван Иванович}}, title={{Длинное название публикации {number}}}, "
            f"year={{2000}}, journal={{Название научного журнала}}, pages={{1--20}}, keywords={{vak}}}}"
            for number in range(1, 41)
        ),
        encoding="utf-8",
    )
    markdown = process_assets(
        "{{list:publications}}", tmp_path, publications_bibliography=bibliography,
    )
    document = Document()
    for line in markdown.splitlines():
        if line.startswith("[[PUBLICATION:"):
            document.add_paragraph(line)
    _format_publication_paragraphs(document, {"publications": {"hanging_indent": "7.5mm"}})

    assert len(document.paragraphs) == 40
    assert all(paragraph._p.xpath("./w:pPr/w:numPr/w:numId") for paragraph in document.paragraphs)
    assert "[[PUBLICATION:" not in "\n".join(paragraph.text for paragraph in document.paragraphs)


def test_abstract_objects_and_cross_references_share_one_stress_scenario(tmp_path: Path) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    (assets / "data.csv").write_text("Показатель,Значение\nТочность,0.93\n", encoding="utf-8")
    (assets / "scheme.svg").write_text("<svg xmlns='http://www.w3.org/2000/svg'/>", encoding="utf-8")
    markdown = """См. {{ref:table:data}}, {{ref:figure:scheme}} и {{ref:equation:model}}.

```table
source: assets/data.csv
id: data
caption: Результаты
```
```figure
source: assets/scheme.svg
id: scheme
caption: Схема
```
```equation
id: model
latex: Q = A + B
```
"""

    result = process_assets(markdown, tmp_path, object_numbering="global")

    assert "[[REF:table:data:1]]" in result
    assert "[[REF:figure:scheme:1]]" in result
    assert "[[REF:equation:model:1]]" in result
    assert "[[TARGET:table:data:0:1]]" in result
    assert "[[TARGET:figure:scheme:0:1]]" in result
    assert "[[EQUATION:model:0:1]]" in result
