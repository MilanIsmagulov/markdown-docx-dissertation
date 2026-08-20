from pathlib import Path

import yaml

from builder.scaffold import scaffold

from builder.parser import parse_note
from builder.resolver import build_index
from builder.validator import validate_index
from builder.assembler import assemble_note
from builder.assets import process_assets
from builder import assets as assets_module
from builder.docx_renderer import (
    _apply_section_profiles,
    _format_numbered_equations,
    _format_publication_paragraphs,
    _install_cross_reference_fields,
    _prepend_abstract_front_matter,
    _set_font,
)
from builder.config import load_document_config
from builder.cli import archive_previous_versions, next_versioned_output
from builder.bibliography import BibEntry, format_publication, format_publication_authors, publication_counts, read_bib_entries
from builder.statistics import measure_docx_pages, read_statistics_manifest, source_statistics, write_statistics_manifest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm
from pypdf import PdfWriter


def write_note(root: Path, relative: str, text: str) -> Path:
    path = root / "content" / relative
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    return path


def write_abstract_validation_config(root: Path, *, denylist: list[str] | None = None) -> None:
    config = root / "config"
    config.mkdir(exist_ok=True)
    (config / "metadata.yaml").write_text(
        yaml.safe_dump(
            {
                "title": "Title", "degree": "Degree", "city": "City", "year": 2026,
                "author": {"full_name": "Author"},
                "organization": {"full_name": "Organization"},
                "specialty": {"code": "2.3.1", "name": "Specialty"},
                "supervisor": {"full_name": "Supervisor", "degree": "Degree", "title": "Title"},
            },
            allow_unicode=True,
        ),
        encoding="utf-8",
    )
    (config / "abstract.yaml").write_text(
        yaml.safe_dump(
            {
                "validation": {
                    "public_profile": True,
                    "placeholder_markers": ["[specify", "_____", "20__"],
                    "allowed_placeholder_fields": [
                        "abstract.defense.date", "abstract.defense.time",
                        "abstract.defense.mailing_date",
                    ],
                    "denylist": denylist or [],
                },
                "defense": {
                    "organization": "Organization", "organization_unit": "Unit",
                    "opponents": [{"degree": "Degree", "full_name": "Opponent"}],
                    "leading_organization": "[specify organization]", "date": "_____ 20__",
                    "time": "_____", "council": "Council", "address": "Address",
                    "library": "Library", "website": "https://example.test",
                    "mailing_date": "_____ 20__", "secretary": "Secretary",
                    "secretary_degree": "Degree",
                },
            },
            allow_unicode=True,
        ),
        encoding="utf-8",
    )
    (config / "research.yaml").write_text(
        "chapters:\n  - {id: 'chapter:1', summary: 'abstract:chapter:1'}\n",
        encoding="utf-8",
    )


def test_document_config_resolves_bibliography_and_csl(tmp_path: Path) -> None:
    (tmp_path / "config").mkdir()
    (tmp_path / "bibliography").mkdir()
    (tmp_path / "assets" / "csl").mkdir(parents=True)
    (tmp_path / "bibliography" / "library.bib").write_text("", encoding="utf-8")
    (tmp_path / "assets" / "csl" / "gost.csl").write_text("<style/>", encoding="utf-8")
    (tmp_path / "config" / "document.yaml").write_text(
        """document:
  bibliography: bibliography/library.bib
  csl: assets/csl/gost.csl
  bibliography_title: СПИСОК ИСТОЧНИКОВ
""",
        encoding="utf-8",
    )

    config = load_document_config(tmp_path)

    assert config.bibliography == tmp_path / "bibliography" / "library.bib"
    assert config.csl == tmp_path / "assets" / "csl" / "gost.csl"
    assert config.bibliography_title == "СПИСОК ИСТОЧНИКОВ"


def test_build_output_is_versioned_and_previous_versions_are_archived(tmp_path: Path) -> None:
    build = tmp_path / "build"
    old = build / ".old"
    old.mkdir(parents=True)
    (build / "dissertation-v10.docx").write_bytes(b"ten")
    (old / "dissertation-v11.docx").write_bytes(b"eleven")

    output = next_versioned_output(build / "dissertation-v10.docx")
    assert output.name == "dissertation-v12.docx"
    output.write_bytes(b"twelve")

    archived = archive_previous_versions(output)
    assert [path.name for path in archived] == ["dissertation-v10.docx"]
    assert not (build / "dissertation-v10.docx").exists()
    assert (old / "dissertation-v10.docx").read_bytes() == b"ten"
    assert output.read_bytes() == b"twelve"


def test_parser_reads_metadata_links_and_transclusions(tmp_path: Path) -> None:
    path = write_note(
        tmp_path,
        "root.md",
        "---\nid: document:root\ntitle: Root\n---\n# Intro\n[[eq:loss|formula]]\n![[Method#Details]]\n",
    )
    note = parse_note(path)
    assert note.note_id == "document:root"
    assert [link.kind for link in note.links] == ["link", "transclusion"]
    assert note.links[1].heading == "Details"


def test_parser_ignores_links_inside_code(tmp_path: Path) -> None:
    path = write_note(tmp_path, "root.md", "`[[inline]]`\n```md\n![[fenced]]\n```\n[[real]]\n")
    note = parse_note(path)
    assert [link.target for link in note.links] == ["real"]


def test_resolver_accepts_stable_id_and_heading(tmp_path: Path) -> None:
    root = write_note(tmp_path, "root.md", "# Root\n![[method:one#Details]]\n")
    write_note(tmp_path, "methods/renamed.md", "---\nid: method:one\n---\n# Method\n## Details\nText\n")
    index = build_index(tmp_path, tmp_path / "content")
    assert validate_index(index, root) == []


def test_validator_reports_missing_heading_and_target(tmp_path: Path) -> None:
    root = write_note(tmp_path, "root.md", "# Root\n[[Missing]]\n![[Existing#Absent]]\n")
    write_note(tmp_path, "Existing.md", "# Existing\n## Present\n")
    diagnostics = validate_index(build_index(tmp_path, tmp_path / "content"), root)
    assert {item.code for item in diagnostics} == {"E_LINK_MISSING", "E_HEADING_MISSING"}


def test_validator_reports_transclusion_cycle(tmp_path: Path) -> None:
    root = write_note(tmp_path, "A.md", "# A\n![[B]]\n")
    write_note(tmp_path, "B.md", "# B\n![[A]]\n")
    diagnostics = validate_index(build_index(tmp_path, tmp_path / "content"), root)
    assert "E_TRANSCLUSION_CYCLE" in {item.code for item in diagnostics}


def test_assembler_expands_equation_note_as_markdown_math(tmp_path: Path) -> None:
    root = write_note(tmp_path, "root.md", "# Root\n![[equation:loss]]\n")
    write_note(
        tmp_path,
        "equations/loss.md",
        "---\nid: equation:loss\ntype: equation\n---\n$$\nL(\\theta) = -\\sum_i y_i \\log p_i\n$$\n",
    )
    index = build_index(tmp_path, tmp_path / "content")
    root_note = next(note for note in index.notes if note.path == root)
    assembled = assemble_note(index, root_note)
    assert "![[" not in assembled
    assert "$$\nL(\\theta) = -\\sum_i y_i \\log p_i\n$$" in assembled


def test_word_heading_font_does_not_keep_theme_reference() -> None:
    document = Document()
    style = document.styles["Heading 1"]
    _set_font(style, "Times New Roman", "14pt", True)
    fonts = style.element.get_or_add_rPr().rFonts
    assert fonts.get(qn("w:ascii")) == "Times New Roman"
    assert fonts.get(qn("w:hAnsi")) == "Times New Roman"
    assert fonts.get(qn("w:eastAsia")) == "Times New Roman"
    assert fonts.get(qn("w:cs")) == "Times New Roman"
    assert fonts.get(qn("w:asciiTheme")) is None
    assert fonts.get(qn("w:hAnsiTheme")) is None


def test_word_cross_references_use_seq_ref_fields_and_bookmarks() -> None:
    document = Document()
    document.add_paragraph("См. [[REF:equation:quality:1.1]].")
    document.add_paragraph("([[TARGET:equation:quality:1:1]])")

    _install_cross_reference_fields(document)

    xml = document._element.xml
    assert "SEQ MdEquationChapter1" in xml
    assert "REF md_dis_equ_" in xml
    assert "w:bookmarkStart" in xml
    assert "w:bookmarkEnd" in xml
    assert "[[REF:" not in xml
    assert "[[TARGET:" not in xml
    for field_run in document._element.xpath(".//w:r[w:instrText]"):
        assert not field_run.xpath("./w:rPr/w:i")
        assert not field_run.xpath("./w:rPr/w:iCs")


def test_word_object_lists_use_ref_and_pageref_fields() -> None:
    document = Document()
    document.add_paragraph("[[LIST:figure]]")
    document.add_paragraph("Всего [[STAT:pages]] страниц")
    document.add_paragraph("Рисунок [[TARGET:figure:pipeline:1:1]] – Схема обработки")

    _install_cross_reference_fields(document)

    xml = document._element.xml
    assert "[[LIST:" not in xml
    assert "REF md_dis_fig_" in xml
    assert "PAGEREF md_dis_fig_" in xml
    assert "NUMPAGES" in xml
    assert "Схема обработки" in xml


def test_asset_processor_numbers_csv_figure_equation_and_references(tmp_path: Path) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    (assets / "data.csv").write_text("Показатель,Значение\nТочность,0.93\n", encoding="utf-8")
    (assets / "scheme.svg").write_text("<svg xmlns='http://www.w3.org/2000/svg'/>", encoding="utf-8")
    markdown = """# ВВЕДЕНИЕ

Вводный текст без нумерации главы.

# Глава 1

См. {{ref:table:data}}, {{ref:figure:scheme}} и {{ref:equation:score}}.

```table
source: assets/data.csv
id: data
caption: Результаты
```

```figure
source: assets/scheme.svg
id: scheme
caption: Схема
```

```equation
id: score
latex: Q = A + B
```
"""
    result = process_assets(markdown, tmp_path)
    assert "[[REF:table:data:1.1]], [[REF:figure:scheme:1.1]] и [[REF:equation:score:1.1]]" in result
    assert "Таблица [[TARGET:table:data:1:1]] – Результаты" in result
    assert "| Показатель | Значение |" in result
    assert "![Рисунок [[TARGET:figure:scheme:1:1]] – Схема]" in result
    assert "[[EQUATION:score:1:1]]" in result
    assert "$$\nQ = A + B\n$$" in result


def test_abstract_object_numbering_is_global_across_chapter_headings(tmp_path: Path) -> None:
    markdown = """# Глава 1

```equation
id: shared-model
latex: a = b
```

# Глава 2

См. {{ref:equation:shared-model}} и {{ref:equation:second-model}}.

```equation
id: second-model
latex: c = d
```
"""

    dissertation = process_assets(markdown, tmp_path, object_numbering="chapter")
    abstract = process_assets(markdown, tmp_path, object_numbering="global")

    assert "[[EQUATION:shared-model:1:1]]" in dissertation
    assert "[[EQUATION:second-model:2:1]]" in dissertation
    assert "[[REF:equation:second-model:2.1]]" in dissertation
    assert "[[EQUATION:shared-model:0:1]]" in abstract
    assert "[[EQUATION:second-model:0:2]]" in abstract
    assert "[[REF:equation:shared-model:1]]" in abstract
    assert "[[REF:equation:second-model:2]]" in abstract


def test_abstract_and_dissertation_use_document_scoped_bookmarks() -> None:
    dissertation = Document()
    dissertation.add_paragraph("([[TARGET:equation:shared-model:1:1]])")
    dissertation.add_paragraph("См. [[REF:equation:shared-model:1.1]].")
    abstract = Document()
    abstract.add_paragraph("([[TARGET:equation:shared-model:0:1]])")
    abstract.add_paragraph("См. [[REF:equation:shared-model:1]].")

    _install_cross_reference_fields(dissertation, "dissertation")
    _install_cross_reference_fields(abstract, "abstract")

    dissertation_xml = dissertation._element.xml
    abstract_xml = abstract._element.xml
    assert "md_dis_equ_" in dissertation_xml
    assert "MdEquationChapter1" in dissertation_xml
    assert "md_abs_equ_" in abstract_xml
    assert "MdAbstractEquation" in abstract_xml
    assert "MdEquationChapter" not in abstract_xml


def test_abstract_equation_tabs_fit_the_a5_body_width() -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(148)
    section.left_margin = Mm(15)
    section.right_margin = Mm(15)
    document.add_paragraph("[[EQUATION:model:0:1]]")
    formula = document.add_paragraph()
    math_para = OxmlElement("m:oMathPara")
    math_para.append(OxmlElement("m:oMath"))
    formula._p.append(math_para)

    _format_numbered_equations(document, {"body": {"font": {"family": "Times New Roman"}}})

    tab_positions = formula._p.xpath("./w:pPr/w:tabs/w:tab/@w:pos")
    expected_width = int(
        section.page_width.twips - section.left_margin.twips - section.right_margin.twips
    )
    assert tab_positions == [str(expected_width // 2), str(expected_width)]


def test_validator_scopes_duplicate_object_ids_to_selected_document(tmp_path: Path) -> None:
    content = tmp_path / "content"
    content.mkdir()
    dissertation_root = write_note(tmp_path, "content/dissertation.md", "![[chapter]]\n")
    abstract_root = write_note(
        tmp_path,
        "content/abstract.md",
        "---\ntype: abstract\n---\n\n![[summary]]\n",
    )
    write_note(
        tmp_path,
        "content/chapter.md",
        "```equation\nid: shared-model\nlatex: a = b\n```\n\n{{ref:equation:shared-model}}\n",
    )
    write_note(
        tmp_path,
        "content/summary.md",
        "```equation\nid: shared-model\nlatex: c = d\n```\n\n{{ref:equation:shared-model}}\n",
    )
    index = build_index(tmp_path, content)

    dissertation_diagnostics = validate_index(index, dissertation_root)
    abstract_diagnostics = validate_index(index, abstract_root)

    assert not any(item.code == "E_OBJECT_ID_DUPLICATE" for item in dissertation_diagnostics)
    assert not any(item.code == "E_OBJECT_ID_DUPLICATE" for item in abstract_diagnostics)


def test_validator_does_not_resolve_object_reference_from_another_document(tmp_path: Path) -> None:
    content = tmp_path / "content"
    content.mkdir()
    dissertation_root = write_note(tmp_path, "content/dissertation.md", "![[chapter]]\n")
    abstract_root = write_note(
        tmp_path,
        "content/abstract.md",
        "---\ntype: abstract\n---\n\n![[summary]]\n",
    )
    write_note(
        tmp_path,
        "content/chapter.md",
        "```equation\nid: dissertation-only\nlatex: a = b\n```\n",
    )
    write_note(
        tmp_path,
        "content/summary.md",
        "См. {{ref:equation:dissertation-only}}.\n",
    )
    index = build_index(tmp_path, content)

    dissertation_diagnostics = validate_index(index, dissertation_root)
    abstract_diagnostics = validate_index(index, abstract_root)

    assert not any(item.code == "E_OBJECT_REFERENCE_MISSING" for item in dissertation_diagnostics)
    assert any(item.code == "E_OBJECT_REFERENCE_MISSING" for item in abstract_diagnostics)


def test_abstract_validator_allows_placeholders_only_in_draft_mode(tmp_path: Path) -> None:
    write_abstract_validation_config(tmp_path)
    root = write_note(
        tmp_path, "abstract.md",
        "---\nid: document:abstract\ntype: abstract\n---\n\n![[summary]]\n",
    )
    write_note(
        tmp_path, "summary.md",
        "---\nid: abstract:chapter:1\ntype: chapter-summary\nsource: chapter:1\n---\n\nSummary\n",
    )
    index = build_index(tmp_path, tmp_path / "content")

    draft = validate_index(index, root, validation_mode="draft")
    final = validate_index(index, root, validation_mode="final")

    assert any(item.code == "W_ABSTRACT_PLACEHOLDER" and item.severity == "warning" for item in draft)
    assert not any(item.code == "E_ABSTRACT_PLACEHOLDER" for item in draft)
    assert any(item.code == "E_ABSTRACT_PLACEHOLDER" and item.severity == "error" for item in final)
    assert not any(
        item.code == "E_ABSTRACT_PLACEHOLDER"
        and item.message.endswith(("abstract.defense.date", "abstract.defense.time", "abstract.defense.mailing_date"))
        for item in final
    )


def test_abstract_validator_checks_chapter_summary_count_and_unique_source(tmp_path: Path) -> None:
    write_abstract_validation_config(tmp_path)
    root = write_note(
        tmp_path, "abstract.md",
        "---\ntype: abstract\n---\n\n![[summary-a]]\n![[summary-b]]\n",
    )
    for name in ("summary-a", "summary-b"):
        write_note(
            tmp_path, f"{name}.md",
            f"---\nid: abstract:{name}\ntype: chapter-summary\nsource: chapter:1\n---\n\nSummary\n",
        )

    diagnostics = validate_index(build_index(tmp_path, tmp_path / "content"), root)
    codes = {item.code for item in diagnostics}

    assert "E_ABSTRACT_CHAPTER_SOURCE_DUPLICATE" in codes
    assert "E_ABSTRACT_CHAPTER_COUNT" in codes


def test_abstract_validator_applies_configured_public_denylist(tmp_path: Path) -> None:
    write_abstract_validation_config(tmp_path, denylist=["private-token"])
    root = write_note(
        tmp_path, "abstract.md",
        "---\ntype: abstract\n---\n\n![[summary]]\n\nprivate-token\n",
    )
    write_note(
        tmp_path, "summary.md",
        "---\nid: abstract:chapter:1\ntype: chapter-summary\nsource: chapter:1\n---\n\nSummary\n",
    )

    diagnostics = validate_index(build_index(tmp_path, tmp_path / "content"), root)

    assert any(item.code == "E_PUBLIC_DENYLIST" for item in diagnostics)


def test_asset_processor_expands_object_and_term_lists(tmp_path: Path) -> None:
    content = tmp_path / "content"
    content.mkdir()
    config = tmp_path / "config"
    config.mkdir()
    (config / "terms.yaml").write_text(
        "abbreviations:\n  - term: ИИ\n    definition: искусственный интеллект\nsymbols: []\n",
        encoding="utf-8",
    )
    result = process_assets(
        "{{list:figures}}\n\n{{list:tables}}\n\n{{list:abbreviations}}\n\n{{list:symbols}}\n",
        content,
    )
    assert "[[LIST:figure]]" in result
    assert "[[LIST:table]]" in result
    assert "| ИИ | искусственный интеллект |" in result


def test_document_and_publication_statistics_are_expanded(tmp_path: Path) -> None:
    content = tmp_path / "content"
    content.mkdir()
    main_bib = tmp_path / "main.bib"
    main_bib.write_text("@book{one, title={One}, author={A}, year={2026}, publisher={P}}", encoding="utf-8")
    publications = tmp_path / "publications.bib"
    publications.write_text(
        "@article{pub, author={A}, title={T}, journal={J}, year={2026}, keywords={vak, scopus}}\n"
        "@software{soft, author={A}, title={S}, number={1}, year={2026}, keywords={software-registration}}",
        encoding="utf-8",
    )
    result = process_assets(
        "# Глава 1\n# ПРИЛОЖЕНИЕ А\n{{stat:pages}}/{{stat:chapters}}/{{stat:appendices}}/"
        "{{stat:bibliography}}/{{stat:publications}}/{{stat:publications_vak}}/"
        "{{stat:publications_scopus_wos}}/{{stat:software_registrations}}/{{stat_phrase:figures}}/"
        "{{stat_phrase:appendices}}",
        content,
        main_bib,
        publications,
    )
    assert "[[STAT:pages]]/1/1/1/2/1/1/1/0 рисунков/1 приложение" in result
    counts = publication_counts(read_bib_entries(publications))
    assert counts["publications"] == 2


def test_statistics_override_materializes_dissertation_pages_and_inflection(tmp_path: Path) -> None:
    result = process_assets(
        "{{stat:pages}}; {{stat_phrase:pages}}; включая {{stat_phrase:tables:acc}}; "
        "на {{stat_phrase:conferences:prep}}",
        tmp_path,
        statistics_override={"pages": 176, "tables": 1, "conferences": 2},
    )

    assert result.strip() == "176; 176 страниц; включая 1 таблицу; на 2 мероприятиях"


def test_pending_publications_are_excluded_from_all_counters(tmp_path: Path) -> None:
    publications = tmp_path / "publications.bib"
    publications.write_text(
        "@article{done, author={A}, title={T}, journal={J}, year={2026}, keywords={vak, scopus}}\n"
        "@article{pending, title={T2}, status={pending}, keywords={vak, wos}}",
        encoding="utf-8",
    )

    counts = publication_counts(read_bib_entries(publications))

    assert counts["publications"] == 1
    assert counts["publications_vak"] == 1
    assert counts["publications_scopus"] == 1
    assert counts["publications_wos"] == 0

    root = write_note(tmp_path, "root.md", "# Root\n")
    diagnostics = validate_index(
        build_index(tmp_path, tmp_path / "content"), root,
        publications_bibliography=publications,
    )
    assert not any("pending" in item.message and item.code == "E_BIB_FIELD_MISSING" for item in diagnostics)


def test_page_count_uses_authoritative_pdf_and_statistics_manifest(tmp_path: Path) -> None:
    docx = tmp_path / "build" / "dissertation-v1.docx"
    docx.parent.mkdir()
    docx.write_bytes(b"docx-placeholder")
    pdf = tmp_path / "dissertation.pdf"
    writer = PdfWriter()
    for _ in range(3):
        writer.add_blank_page(width=595, height=842)
    with pdf.open("wb") as stream:
        writer.write(stream)

    pages = measure_docx_pages(docx, pdf)
    manifest = write_statistics_manifest(tmp_path, docx, {"pages": pages, "figures": 2})
    source, statistics = read_statistics_manifest(tmp_path)

    assert manifest.is_file()
    assert source == docx
    assert statistics == {"pages": 3, "figures": 2}


def test_source_statistics_excludes_abstract_only_objects(tmp_path: Path) -> None:
    write_note(tmp_path, "chapter.md", "---\ntype: chapter\n---\n\n```figure\nid: f\n```\n")
    write_note(tmp_path, "appendix.md", "---\ntype: appendix\n---\n\n```table\nid: t\n```\n")
    write_note(
        tmp_path, "abstract.md",
        "---\ntype: abstract\n---\n\n```figure\nid: abstract-only\n```\n",
    )

    statistics = source_statistics(
        build_index(tmp_path, tmp_path / "content"), None, None, None,
    )

    assert statistics["chapters"] == 1
    assert statistics["appendices"] == 1
    assert statistics["figures"] == 1
    assert statistics["tables"] == 1


def test_conference_registry_generates_list_and_statistics(tmp_path: Path) -> None:
    conferences = tmp_path / "conferences.bib"
    conferences.write_text(
        "@conference{event, title={Conference}, eventdate={2026-02-20}, "
        "location={City}, talk={Report}, keywords={national}}",
        encoding="utf-8",
    )

    result = process_assets(
        "{{stat:conferences}} {{stat_phrase:conferences}}\n\n{{list:conferences}}",
        tmp_path,
        conferences_bibliography=conferences,
    )

    assert "1 1 мероприятие" in result
    assert "- Conference (2026-02-20, City); доклад «Report»." in result


def test_publication_list_uses_yaml_groups_initials_and_markers(tmp_path: Path) -> None:
    publications = tmp_path / "publications.bib"
    publications.write_text(
        "@online{site, author={Иван Иванович Иванов and Петров, Петр Петрович}, "
        "title={Ресурс}, year={2026}, url={https://example.test/a-b}, urldate={2026-08-20}, keywords={web}}",
        encoding="utf-8",
    )
    result = process_assets(
        "{{list:publications}}", tmp_path, publications_bibliography=publications,
        publication_config={"groups": [{"id": "web", "title": "Сетевые работы", "keywords_any": ["web"]}]},
    )

    assert "**Сетевые работы:**" in result
    assert "[[PUBLICATION:1]] Иванов И. И.; Петров П. П." in result
    assert "[Электронный ресурс]" in result
    assert "URL: https://example.test/a-b" in result
    assert "дата обращения: 2026-08-20" in result


def test_publication_author_order_and_entry_types_are_preserved() -> None:
    assert format_publication_authors("Иванов, Иван Иванович and Петр Петрович Петров") == (
        "Иванов И. И.; Петров П. П."
    )
    entry = BibEntry("software", "certificate", {
        "author": "Иванов, Иван Иванович", "title": "Программа", "year": "2026", "number": "2000123456",
    })
    assert "Свидетельство о государственной регистрации программы для ЭВМ 2000123456" in format_publication(entry)


def test_publication_docx_paragraph_has_hanging_indent() -> None:
    document = Document()
    paragraph = document.add_paragraph("[[PUBLICATION:12]] Long DOI: 10.1000/example")

    _format_publication_paragraphs(document, {"publications": {"hanging_indent": "8mm"}})

    assert paragraph.text.startswith("Long DOI")
    assert paragraph._p.xpath("./w:pPr/w:numPr/w:numId")
    children = list(document.part.numbering_part.element)
    first_num = next(i for i, node in enumerate(children) if node.tag == qn("w:num"))
    assert all(node.tag == qn("w:abstractNum") for node in children[:first_num])


def test_validator_checks_object_references_assets_and_citations(tmp_path: Path) -> None:
    root = write_note(
        tmp_path,
        "root.md",
        """# Root

См. {{number:figure:missing}} и [@absent].

```figure
id: scheme
caption: Схема
source: assets/missing.png
```
""",
    )
    bibliography = tmp_path / "bibliography.bib"
    bibliography.write_text("@article{unused, title={Unused}}\n", encoding="utf-8")
    diagnostics = validate_index(build_index(tmp_path, tmp_path / "content"), root, bibliography)
    codes = {item.code for item in diagnostics}
    assert "E_ASSET_MISSING" in codes
    assert "E_OBJECT_REFERENCE_MISSING" in codes
    assert "E_CITATION_MISSING" in codes
    assert "W_BIB_UNUSED" in codes


def test_validator_detects_duplicate_bibliographic_records(tmp_path: Path) -> None:
    root = write_note(tmp_path, "root.md", "# Root\n")
    bibliography = tmp_path / "publications.bib"
    bibliography.write_text(
        "@article{one, author={Иванов, И. И.}, title={Одинаковая работа}, year={2026}, journal={Журнал}}\n"
        "@inproceedings{two, author={Иванов, И. И.}, title={Одинаковая работа}, year={2026}, booktitle={Труды}}",
        encoding="utf-8",
    )

    diagnostics = validate_index(
        build_index(tmp_path, tmp_path / "content"), root, publications_bibliography=bibliography,
    )

    assert "E_BIB_RECORD_DUPLICATE" in {item.code for item in diagnostics}


def test_pdf_transclusion_is_validated_preserved_and_expanded(tmp_path: Path, monkeypatch) -> None:
    root = write_note(tmp_path, "root.md", "![[appendix]]\n")
    write_note(tmp_path, "appendix.md", "# ПРИЛОЖЕНИЕ А\n\n![[assets/document.pdf]]\n")
    pdf = tmp_path / "content" / "assets" / "document.pdf"
    pdf.parent.mkdir(parents=True)
    pdf.write_bytes(b"%PDF-demo")
    index = build_index(tmp_path, tmp_path / "content")
    assert validate_index(index, root) == []
    root_note = next(note for note in index.notes if note.path == root)
    assembled = assemble_note(index, root_note)
    assert "![[assets/document.pdf]]" in assembled

    page_one = tmp_path / "page-1.png"
    page_two = tmp_path / "page-2.png"
    page_one.write_bytes(b"png")
    page_two.write_bytes(b"png")
    monkeypatch.setattr(assets_module, "_render_pdf_pages", lambda source, content, page=None: [page_one, page_two])
    expanded = process_assets(assembled, tmp_path / "content")
    assert "page-1.png" in expanded
    assert "[[PDF_PAGE_BREAK]]" in expanded
    assert "page-2.png" in expanded
    assert "{width=175mm}" in expanded


def test_section_profile_marker_is_preserved_for_docx_renderer(tmp_path: Path) -> None:
    result = process_assets("{{section:appendix-wide}}\n\n# Appendix", tmp_path)

    assert "[[SECTION:appendix-wide]]" in result


def test_validator_reports_unknown_section_profile(tmp_path: Path) -> None:
    root = write_note(tmp_path, "root.md", "{{section:missing}}\n")
    sections = tmp_path / "config" / "sections.yaml"
    sections.parent.mkdir(parents=True)
    sections.write_text(
        "default: dissertation\nprofiles:\n  dissertation:\n    size: A4\n"
        "    orientation: portrait\n    margins: {left: 25mm, right: 10mm, top: 20mm, bottom: 20mm}\n",
        encoding="utf-8",
    )

    diagnostics = validate_index(build_index(tmp_path, tmp_path / "content"), root)

    assert any(item.code == "E_SECTION_UNKNOWN" for item in diagnostics)


def test_landscape_section_profile_creates_real_word_section(tmp_path: Path) -> None:
    sections = tmp_path / "config" / "sections.yaml"
    sections.parent.mkdir(parents=True)
    sections.write_text(
        "default: dissertation\nprofiles:\n"
        "  dissertation:\n    margins: {left: 25mm, right: 10mm, top: 20mm, bottom: 20mm}\n"
        "  landscape:\n    orientation: landscape\n"
        "    margins: {left: 20mm, right: 10mm, top: 15mm, bottom: 15mm}\n",
        encoding="utf-8",
    )
    document = Document()
    document.add_paragraph("Portrait")
    document.add_paragraph("[[SECTION:landscape]]")
    document.add_paragraph("Landscape")

    _apply_section_profiles(document, tmp_path, {"document": {"margins": {}}})

    assert len(document.sections) == 2
    assert round(document.sections[1].page_width / Mm(1)) == 297
    assert round(document.sections[1].page_height / Mm(1)) == 210


def test_validator_reports_missing_semantic_source(tmp_path: Path) -> None:
    root = write_note(tmp_path, "root.md", "![[summary]]\n")
    write_note(
        tmp_path,
        "summary.md",
        "---\nid: abstract:chapter:1\ntype: chapter-summary\nsource: chapter:1\n---\n\nSummary\n",
    )

    diagnostics = validate_index(build_index(tmp_path, tmp_path / "content"), root)

    assert any(item.code == "E_SEMANTIC_SOURCE_MISSING" for item in diagnostics)


def test_abstract_front_matter_has_two_unnumbered_sections_and_restarts_at_one() -> None:
    document = Document()
    document.add_paragraph("Body")
    metadata = {
        "title": "Title",
        "degree": "кандидата технических наук",
        "author": {"full_name": "Author"},
        "specialty": {"code": "2.3.1", "name": "Specialty"},
        "supervisor": {"full_name": "Supervisor", "degree": "д.т.н.", "title": "профессор"},
        "city": "City",
        "year": 2026,
    }

    _prepend_abstract_front_matter(document, metadata, {"defense": {}})

    assert len(document.sections) == 2
    assert not document.sections[0]._sectPr.xpath("./w:headerReference")
    assert document.sections[1]._sectPr.xpath("./w:pgNumType/@w:start") == ["1"]


def test_abstract_front_matter_uses_configured_font_sizes() -> None:
    document = Document()
    document.add_paragraph("Body")
    metadata = {
        "title": "Title", "degree": "Degree", "author": {"full_name": "Author"},
        "specialty": {"code": "2.3.1", "name": "Specialty"},
        "supervisor": {"full_name": "Supervisor", "degree": "Degree", "title": "Title"},
        "city": "City", "year": 2026,
    }
    config = {
        "layout": {"title_font_size": "11pt", "verso_font_size": "10pt"},
        "defense": {"opponents": [{"degree": "Degree", "full_name": "ИВАНОВ Иван Иванович"}]},
    }

    _prepend_abstract_front_matter(document, metadata, config)

    assert document.paragraphs[0].runs[0].font.size.pt == 11
    table = document.tables[0]
    verso_runs = [
        run
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
        for run in paragraph.runs
    ]
    assert verso_runs
    assert all(run.font.size is None or run.font.size.pt == 10 for run in verso_runs)
    assert any(run.text == "SUPERVISOR" and run.bold for run in verso_runs)
    assert any(run.text.endswith("ИВАНОВ Иван Иванович") and run.bold for run in verso_runs)
    assert len(document.tables) == 1
    assert len(table.columns) == 2
    assert table.autofit is True
    assert table._tbl.tblPr.xpath("./w:tblW/@w:type") == ["pct"]
    assert table._tbl.tblPr.xpath("./w:tblW/@w:w") == ["5000"]
    assert table._tbl.tblPr.xpath("./w:tblCaption/@w:val") == ["abstract-layout"]
    assert "Научный руководитель:" in "\n".join(cell.text for cell in table.columns[0].cells)
    assert "Официальные оппоненты:" in "\n".join(cell.text for cell in table.columns[0].cells)
    assert any(
        run.text.endswith("ИВАНОВ Иван Иванович") and run.bold
        for row in table.rows for cell in row.cells for paragraph in cell.paragraphs for run in paragraph.runs
    )
    assert all(run.font.size.pt == 10 for run in verso_runs if run.text)
    assert table.rows[0].cells[0].paragraphs[0].paragraph_format.space_after.pt == 30
    spaced_opponent_names = [
        paragraph
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
        if paragraph.text == "ИВАНОВ Иван Иванович"
        and paragraph.paragraph_format.space_after is not None
        and paragraph.paragraph_format.space_after.pt == 40
    ]
    assert len(spaced_opponent_names) == 1


def test_validator_reports_missing_pdf_asset(tmp_path: Path) -> None:
    root = write_note(tmp_path, "root.md", "![[assets/missing.pdf]]\n")
    diagnostics = validate_index(build_index(tmp_path, tmp_path / "content"), root)
    assert {item.code for item in diagnostics} == {"E_ASSET_MISSING"}


def test_scaffold_creates_nested_obsidian_structure_and_is_idempotent(tmp_path: Path) -> None:
    config = tmp_path / "config"
    config.mkdir()
    (config / "metadata.yaml").write_text(
        yaml.safe_dump(
            {"structure": {"chapters": 2, "paragraphs_per_chapter": 2}},
            allow_unicode=True,
        ),
        encoding="utf-8",
    )
    existing = tmp_path / "content" / "01 Chapter" / "1.1. Авторское название.md"
    existing.parent.mkdir(parents=True)
    existing.write_text("## 1.1. Авторское название\n\nАвторский текст.\n", encoding="utf-8")

    created, chapters = scaffold(tmp_path)
    assert chapters == 2
    assert created == 28
    root = (tmp_path / "content" / "root.md").read_text(encoding="utf-8")
    chapter = (tmp_path / "content" / "01 Chapter" / "Глава 1. Тест.md").read_text(encoding="utf-8")
    assert "![[01 Chapter/Глава 1. Тест]]" in root
    assert "![[00 Front Matter/Введение]]" in root
    assert "![[00 Front Matter/Списки]]" in root
    assert "![[90 Back Matter/Заключение]]" in root
    assert "![[90 Back Matter/Список литературы]]" in root
    assert "![[99 Appendices/Приложения]]" in root
    introduction = (tmp_path / "content" / "00 Front Matter" / "Введение.md").read_text(encoding="utf-8")
    assert "# ВВЕДЕНИЕ" in introduction
    assert "![[00 Front Matter/Актуальность темы]]" in introduction
    assert "![[00 Front Matter/Положения, выносимые на защиту]]" in introduction
    conclusion = (tmp_path / "content" / "90 Back Matter" / "Заключение.md").read_text(encoding="utf-8")
    assert "# ЗАКЛЮЧЕНИЕ" in conclusion
    appendix_index = (tmp_path / "content" / "99 Appendices" / "Приложения.md").read_text(encoding="utf-8")
    assert "![[99 Appendices/Приложение А]]" in appendix_index
    appendix = (tmp_path / "content" / "99 Appendices" / "Приложение А.md").read_text(encoding="utf-8")
    assert "# ПРИЛОЖЕНИЕ А" in appendix
    assert "![[01 Chapter/Преамбула главы 1]]" in chapter
    assert "![[01 Chapter/1.1. Авторское название]]" in chapter
    assert "![[01 Chapter/Выводы по главе 1]]" in chapter
    chapter_two = (tmp_path / "content" / "02 Chapter" / "Глава 2. Тест.md").read_text(encoding="utf-8")
    assert "![[02 Chapter/2.1. Параграф]]" in chapter_two
    assert "![[02 Chapter/2.2. Параграф]]" in chapter_two
    paragraph_two = tmp_path / "content" / "02 Chapter" / "2.2. Параграф.md"
    assert "id: paragraph:2.2" in paragraph_two.read_text(encoding="utf-8")

    paragraph = existing
    paragraph.write_text(paragraph.read_text(encoding="utf-8") + "\nАвторский текст.\n", encoding="utf-8")
    created_again, _ = scaffold(tmp_path)
    assert created_again == 0
    assert "Авторский текст." in paragraph.read_text(encoding="utf-8")
