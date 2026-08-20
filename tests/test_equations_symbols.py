from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from builder.assets import _equation_latex, process_assets
from builder.docx_renderer import _format_equation_where_blocks
from builder.resolver import build_index
from builder.validator import validate_index


def test_equation_lines_build_aligned_and_cases_latex() -> None:
    aligned = _equation_latex({"layout": "aligned", "lines": [r"a &= b", r"c &= d"]})
    cases = _equation_latex({"layout": "cases", "lines": [r"x, & x > 0", r"0, & x \le 0"]})

    assert aligned == r"\begin{aligned} a &= b \\ c &= d \end{aligned}"
    assert cases == r"\begin{cases} x, & x > 0 \\ 0, & x \le 0 \end{cases}"


def test_equation_where_uses_canonical_definitions_and_filters_symbol_list(tmp_path: Path) -> None:
    content = tmp_path / "content"
    content.mkdir()
    config = tmp_path / "config"
    config.mkdir()
    (config / "terms.yaml").write_text(
        "symbols:\n  - {term: Q, definition: итоговая оценка}\n  - {term: A, definition: показатель}\n  - {term: Z, definition: не используется}\n",
        encoding="utf-8",
    )
    markdown = """# Глава 1
```equation
id: model
layout: aligned
lines:
  - Q &= A + 1
  - A &= 2
symbols: [Q, A]
where: [Q, A]
```

{{list:symbols}}
"""

    result = process_assets(markdown, content)

    assert r"\begin{aligned} Q &= A + 1 \\ A &= 2 \end{aligned}" in result
    assert "[[WHERE_BEGIN]]" in result and "[[WHERE_END]]" in result
    assert "$Q$ — итоговая оценка" in result
    assert "| Q | итоговая оценка |" in result
    assert "| A | показатель |" in result
    assert "не используется" not in result


def test_where_block_gets_hanging_layout_and_markers_are_removed() -> None:
    document = Document()
    document.add_paragraph("[[WHERE_BEGIN]]")
    document.add_paragraph("где")
    first = document.add_paragraph("Q — итоговая оценка")
    second = document.add_paragraph("A — показатель качества")
    document.add_paragraph("[[WHERE_END]]")

    _format_equation_where_blocks(document, {"body": {"font": {"family": "Times New Roman"}}})

    assert [paragraph.text for paragraph in document.paragraphs] == ["где", first.text, second.text]
    assert document.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert round(first.paragraph_format.left_indent.mm, 1) == 10.0
    assert round(first.paragraph_format.first_line_indent.mm, 1) == -10.0
    assert first.paragraph_format.keep_with_next is True
    assert second.paragraph_format.keep_with_next is False


def test_validator_links_equation_symbols_to_terms_file(tmp_path: Path) -> None:
    content = tmp_path / "content"
    config = tmp_path / "config"
    content.mkdir()
    config.mkdir()
    (config / "terms.yaml").write_text(
        "symbols:\n  - {term: Q, definition: итоговая оценка}\n  - {term: Q, definition: дубликат}\n  - {term: Z, definition: лишнее обозначение}\n",
        encoding="utf-8",
    )
    root = content / "root.md"
    root.write_text(
        """---
id: document:root
type: document
title: Root
---
```equation
id: model
latex: Q = A
symbols: [Q, A]
where: [Q]
```
""",
        encoding="utf-8",
    )

    diagnostics = validate_index(build_index(tmp_path, content), root)
    codes = [item.code for item in diagnostics]

    assert "E_SYMBOL_UNDEFINED" in codes
    assert "E_SYMBOL_DUPLICATE" in codes
    assert "W_SYMBOL_UNUSED" in codes


def test_validator_rejects_invalid_equation_layout_and_where(tmp_path: Path) -> None:
    content = tmp_path / "content"
    content.mkdir()
    root = content / "root.md"
    root.write_text(
        """---
id: document:root
type: document
title: Root
---
```equation
id: broken
layout: matrix
lines: []
where: broken
```
""",
        encoding="utf-8",
    )

    codes = [item.code for item in validate_index(build_index(tmp_path, content), root)]

    assert "E_EQUATION_CONFIG" in codes
    assert "E_EQUATION_WHERE" in codes
