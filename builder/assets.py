from __future__ import annotations

import csv
import base64
import json
import re
import shutil
import subprocess
from dataclasses import dataclass
from pathlib import Path

import yaml

from .bibliography import completed_entries, format_publication, group_publications, publication_counts, read_bib_entries


DIRECTIVE_RE = re.compile(r"^```(?P<kind>table|figure|equation)\s*$")
REFERENCE_RE = re.compile(r"\{\{ref:(?P<kind>table|figure|equation):(?P<id>[A-Za-z0-9_.-]+)}}")
NUMBER_RE = re.compile(r"\{\{number:(?P<kind>table|figure|equation):(?P<id>[A-Za-z0-9_.-]+)}}")
OBJECT_LIST_RE = re.compile(r"\{\{list:(?P<kind>figures|tables)}}")
TERMS_LIST_RE = re.compile(r"\{\{list:(?P<kind>abbreviations|symbols|glossary)}}")
CONFERENCE_LIST_RE = re.compile(r"\{\{list:conferences}}")
PUBLICATION_LIST_RE = re.compile(r"\{\{list:publications}}")
STAT_RE = re.compile(r"\{\{stat:(?P<name>[a-z_]+)}}")
STAT_PHRASE_RE = re.compile(r"\{\{stat_phrase:(?P<name>[a-z_]+)(?::(?P<case>[a-z]+))?}}")
SECTION_RE = re.compile(r"\{\{section:(?P<name>[a-z][a-z0-9_-]*)}}")
CHAPTER_RE = re.compile(r"^#\s+Глава\s+(?P<number>\d+)\b", re.IGNORECASE)
PDF_EMBED_RE = re.compile(
    r"^\s*!\[\[(?P<target>[^]|#]+\.pdf)(?:#page=(?P<page>\d+))?(?:\|(?P<width>\d+(?:\.\d+)?mm))?]]\s*$",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class ObjectNumber:
    kind: str
    number: str


def _range_bounds(cell_range: str | None, rows: list[list[str]]) -> tuple[int, int, int, int]:
    width = max((len(row) for row in rows), default=0)
    if not cell_range:
        return 1, 1, width, len(rows)
    try:
        from openpyxl.utils.cell import range_boundaries

        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"invalid table range: {cell_range}") from exc
    return min_col, min_row, max_col, max_row


def _slice_rows(rows: list[list[str]], cell_range: str | None) -> list[list[str]]:
    min_col, min_row, max_col, max_row = _range_bounds(cell_range, rows)
    return [
        (row + [""] * max_col)[min_col - 1 : max_col]
        for row in rows[min_row - 1 : max_row]
    ]


def _read_rows(path: Path, sheet: str | None, cell_range: str | None = None) -> list[list[str]]:
    if path.suffix.casefold() in {".csv", ".tsv"}:
        delimiter = "\t" if path.suffix.casefold() == ".tsv" else ","
        with path.open("r", encoding="utf-8-sig", newline="") as source:
            rows = [[cell.strip() for cell in row] for row in csv.reader(source, delimiter=delimiter)]
            return _slice_rows(rows, cell_range)
    if path.suffix.casefold() == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("XLSX import requires the 'openpyxl' package") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            worksheet = workbook[sheet] if sheet else workbook.active
        except KeyError as exc:
            workbook.close()
            raise ValueError(f"XLSX sheet does not exist: {sheet}") from exc
        iterator = worksheet[cell_range] if cell_range else worksheet.iter_rows()
        rows = [["" if cell.value is None else str(cell.value) for cell in row] for row in iterator]
        workbook.close()
        return rows
    raise ValueError(f"unsupported table asset: {path.suffix}")


def _table_config_marker(config: dict) -> str:
    supported = {
        key: config[key]
        for key in ("widths", "align", "merges", "repeat_header")
        if key in config
    }
    payload = json.dumps(supported, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    return "[[TABLE_CONFIG:" + base64.urlsafe_b64encode(payload).decode("ascii").rstrip("=") + "]]"


def _table_chunks(rows: list[list[str]], split_rows: int | None) -> list[list[list[str]]]:
    if not split_rows or len(rows) <= split_rows + 1:
        return [rows]
    header, body = rows[0], rows[1:]
    return [[header, *body[start : start + split_rows]] for start in range(0, len(body), split_rows)]


def _escape_cell(value: str) -> str:
    return value.replace("|", "\\|").replace("\n", "<br>")


def _markdown_table(rows: list[list[str]]) -> str:
    if not rows:
        raise ValueError("table asset contains no rows")
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    header = "| " + " | ".join(_escape_cell(cell) for cell in normalized[0]) + " |"
    separator = "| " + " | ".join("---" for _ in range(width)) + " |"
    body = ["| " + " | ".join(_escape_cell(cell) for cell in row) + " |" for row in normalized[1:]]
    return "\n".join((header, separator, *body))


def _load_directive(lines: list[str], start: int) -> tuple[dict, int]:
    closing = next((index for index in range(start + 1, len(lines)) if lines[index].strip() == "```"), None)
    if closing is None:
        raise ValueError(f"asset directive at line {start + 1} is not closed")
    config = yaml.safe_load("\n".join(lines[start + 1 : closing])) or {}
    if not isinstance(config, dict):
        raise ValueError(f"asset directive at line {start + 1} must contain a YAML mapping")
    return config, closing + 1


def _render_pdf_pages(source: Path, content_dir: Path, page: int | None = None) -> list[Path]:
    pdftoppm = shutil.which("pdftoppm")
    if pdftoppm is None:
        raise ValueError("PDF embedding requires Poppler's pdftoppm executable")
    executable = Path(pdftoppm)
    if executable.suffix.casefold() in {".cmd", ".bat"}:
        runtime_root = executable.parents[2]
        native_candidates = list(runtime_root.glob("native/poppler/**/pdftoppm.exe"))
        if native_candidates:
            pdftoppm = str(native_candidates[0])
    cache = content_dir.parent / "build" / "pdf-pages" / source.stem
    cache.mkdir(parents=True, exist_ok=True)
    prefix = cache / "page"
    for stale_page in cache.glob("page-*.png"):
        stale_page.unlink()
    command = [pdftoppm, "-png", "-r", "200"]
    if page is not None:
        command.extend(["-f", str(page), "-l", str(page)])
    command.extend([str(source), str(prefix)])
    try:
        subprocess.run(command, check=True, capture_output=True, text=True)
    except subprocess.CalledProcessError as exc:
        detail = exc.stderr.strip() if exc.stderr else str(exc)
        raise ValueError(f"failed to render PDF {source}: {detail}") from exc
    pages = sorted(cache.glob("page-*.png"), key=lambda path: int(path.stem.rsplit("-", 1)[1]))
    if page is not None:
        pages = [path for path in pages if int(path.stem.rsplit("-", 1)[1]) == page]
    if not pages:
        raise ValueError(f"PDF produced no pages: {source}")
    return pages


def process_assets(
    markdown: str,
    content_dir: Path,
    bibliography: Path | None = None,
    publications_bibliography: Path | None = None,
    conferences_bibliography: Path | None = None,
    statistics_override: dict[str, int] | None = None,
    object_numbering: str = "chapter",
    publication_config: dict | None = None,
) -> str:
    if object_numbering not in {"chapter", "global"}:
        raise ValueError(f"unsupported object numbering profile: {object_numbering}")
    lines = markdown.splitlines()
    counters = {"table": 0, "figure": 0, "equation": 0}
    labels: dict[tuple[str, str], ObjectNumber] = {}
    chapter = 0
    output: list[str] = []
    index = 0

    while index < len(lines):
        line = lines[index]
        pdf_match = PDF_EMBED_RE.match(line)
        if pdf_match is not None:
            source = (content_dir / pdf_match["target"].strip()).resolve()
            if not source.is_relative_to(content_dir.resolve()) or not source.is_file():
                raise ValueError(f"PDF asset not found: {source}")
            requested_page = int(pdf_match["page"]) if pdf_match["page"] else None
            pages = _render_pdf_pages(source, content_dir, requested_page)
            width = pdf_match["width"] or "175mm"
            for page_index, page_path in enumerate(pages):
                if page_index:
                    output.extend(("[[PDF_PAGE_BREAK]]", ""))
                output.extend((f"![](<{page_path.as_posix()}>){{width={width}}}", ""))
            index += 1
            continue
        chapter_match = CHAPTER_RE.match(line)
        if chapter_match is not None:
            if object_numbering == "chapter":
                chapter = int(chapter_match["number"])
                counters = {kind: 0 for kind in counters}
            output.append(line)
            index += 1
            continue
        match = DIRECTIVE_RE.match(line)
        if match is None:
            output.append(line)
            index += 1
            continue

        kind = match["kind"]
        config, index = _load_directive(lines, index)
        object_id = str(config.get("id", "")).strip()
        caption = str(config.get("caption", "")).strip()
        if not object_id:
            raise ValueError(f"{kind} directive requires an id")
        if (kind, object_id) in labels:
            raise ValueError(f"duplicate {kind} id: {object_id}")
        counters[kind] += 1
        marker_chapter = chapter if object_numbering == "chapter" else 0
        number = f"{marker_chapter}.{counters[kind]}" if marker_chapter else str(counters[kind])
        labels[(kind, object_id)] = ObjectNumber(kind, number)

        if kind == "table":
            source = content_dir / str(config.get("source", ""))
            if not source.is_file():
                raise ValueError(f"table asset not found: {source}")
            rows = _read_rows(
                source,
                str(config["sheet"]) if config.get("sheet") else None,
                str(config["range"]) if config.get("range") else None,
            )
            if not rows or not rows[0]:
                raise ValueError("table asset contains no rows in the selected range")
            split_rows_value = config.get("split_rows")
            split_rows = int(split_rows_value) if split_rows_value is not None else None
            chunks = _table_chunks(rows, split_rows)
            widths = config.get("widths", [])
            width_total = 0.0
            if isinstance(widths, list):
                for value in widths:
                    match_width = re.fullmatch(r"(\d+(?:\.\d+)?)mm", str(value).strip(), re.IGNORECASE)
                    if match_width:
                        width_total += float(match_width.group(1))
            section = str(config.get("section", "portrait")).casefold()
            use_landscape = section == "landscape" or (
                section == "auto" and (len(rows[0]) > 6 or width_total > 175)
            )
            if use_landscape:
                output.extend(("[[SECTION:landscape]]", ""))
            for chunk_index, chunk in enumerate(chunks):
                if chunk_index == 0:
                    table_caption = f"Таблица [[TARGET:{kind}:{object_id}:{marker_chapter}:{counters[kind]}]] – {caption}"
                else:
                    table_caption = f"Продолжение таблицы [[NUMBER:{kind}:{object_id}:{number}]]"
                output.extend((table_caption, "", _table_config_marker(config), "", _markdown_table(chunk), ""))
            note = str(config.get("note", "")).strip()
            source_note = str(config.get("source_note", "")).strip()
            if note:
                output.extend((f"Примечание. {note}", ""))
            if source_note:
                output.extend((f"Источник: {source_note}", ""))
            if use_landscape:
                output.extend(("[[SECTION:dissertation]]", ""))
        elif kind == "figure":
            source = content_dir / str(config.get("source", ""))
            if not source.is_file():
                raise ValueError(f"figure asset not found: {source}")
            width = str(config.get("width", "175mm"))
            output.extend(
                (
                    f"![Рисунок [[TARGET:{kind}:{object_id}:{marker_chapter}:{counters[kind]}]] – {caption}]"
                    f"(<{source.as_posix()}>){{width={width}}}",
                    "",
                )
            )
        else:
            latex = str(config.get("latex", "")).strip()
            if not latex:
                raise ValueError("equation directive requires latex")
            output.extend(
                (
                    f"[[EQUATION:{object_id}:{marker_chapter}:{counters[kind]}]]",
                    "",
                    "$$",
                    latex,
                    "$$",
                    "",
                )
            )

    assembled = "\n".join(output)
    assembled = OBJECT_LIST_RE.sub(
        lambda match: f"[[LIST:{'figure' if match['kind'] == 'figures' else 'table'}]]",
        assembled,
    )

    def replace_terms(match: re.Match[str]) -> str:
        terms_path = content_dir.parent / "config" / "terms.yaml"
        if not terms_path.is_file():
            raise ValueError(f"terms configuration not found: {terms_path}")
        data = yaml.safe_load(terms_path.read_text(encoding="utf-8-sig")) or {}
        entries = data.get(match["kind"], [])
        if not isinstance(entries, list):
            raise ValueError(f"terms.{match['kind']} must be a list")
        rows = [["Термин" if match["kind"] == "glossary" else "Обозначение", "Определение" if match["kind"] == "glossary" else "Расшифровка"]]
        for entry in entries:
            if (
                not isinstance(entry, dict)
                or not str(entry.get("term", "")).strip()
                or not str(entry.get("definition", "")).strip()
            ):
                raise ValueError(f"every terms.{match['kind']} entry requires term and definition")
            rows.append([str(entry["term"]).strip(), str(entry["definition"]).strip()])
        return _markdown_table(rows)

    assembled = TERMS_LIST_RE.sub(replace_terms, assembled)
    conferences = sorted(
        completed_entries(read_bib_entries(conferences_bibliography)), key=lambda entry: entry.fields.get("eventdate", "")
    )

    def conference_item(entry) -> str:
        fields = entry.fields
        date = fields.get("eventdate", "").replace("/", " — ")
        details = ", ".join(item for item in (date, fields.get("location", ""), fields.get("organizer", "")) if item)
        text = f"{fields.get('title', entry.key)} ({details})"
        if fields.get("talk"):
            text += f"; доклад «{fields['talk']}»"
        if fields.get("award"):
            text += f"; {fields['award']}"
        return f"- {text}."

    assembled = CONFERENCE_LIST_RE.sub("\n".join(conference_item(entry) for entry in conferences), assembled)
    assembled = SECTION_RE.sub(lambda match: f"[[SECTION:{match['name']}]]", assembled)

    publication_stats = publication_counts(read_bib_entries(publications_bibliography))
    stats = {
        "chapters": len(re.findall(r"^#\s+Глава\s+\d+\b", assembled, re.MULTILINE | re.IGNORECASE)),
        "figures": len(labels_for_kind(labels, "figure")),
        "tables": len(labels_for_kind(labels, "table")),
        "appendices": len(re.findall(r"^#\s+ПРИЛОЖЕНИЕ\s+[А-Я]\b", assembled, re.MULTILINE)),
        "bibliography": len(read_bib_entries(bibliography)),
        "conferences": len(conferences),
        **publication_stats,
    }
    stats.update(statistics_override or {})
    nouns = {
        "chapters": ("глава", "главы", "глав"),
        "figures": ("рисунок", "рисунка", "рисунков"),
        "tables": ("таблица", "таблицы", "таблиц"),
        "appendices": ("приложение", "приложения", "приложений"),
        "bibliography": ("наименование", "наименования", "наименований"),
        "conferences": ("мероприятие", "мероприятия", "мероприятий"),
        "publications": ("печатная работа", "печатные работы", "печатных работ"),
        "publications_vak": ("статья", "статьи", "статей"),
        "publications_scopus_wos": ("публикация", "публикации", "публикаций"),
        "publications_scopus": ("публикация", "публикации", "публикаций"),
        "publications_wos": ("публикация", "публикации", "публикаций"),
        "publications_other": ("работа", "работы", "работ"),
        "patents": ("патент", "патента", "патентов"),
        "software_registrations": ("свидетельство", "свидетельства", "свидетельств"),
        "pages": ("страница", "страницы", "страниц"),
    }
    case_nouns = {
        ("chapters", "acc"): ("главу", "главы", "глав"),
        ("figures", "acc"): ("рисунок", "рисунка", "рисунков"),
        ("tables", "acc"): ("таблицу", "таблицы", "таблиц"),
        ("appendices", "acc"): ("приложение", "приложения", "приложений"),
        ("conferences", "prep"): ("мероприятии", "мероприятиях", "мероприятиях"),
    }

    def inflect(number: int, forms: tuple[str, str, str]) -> str:
        if number % 10 == 1 and number % 100 != 11:
            form = forms[0]
        elif number % 10 in {2, 3, 4} and number % 100 not in {12, 13, 14}:
            form = forms[1]
        else:
            form = forms[2]
        return f"{number} {form}"

    assembled = STAT_PHRASE_RE.sub(
        lambda match: inflect(
            stats[match["name"]],
            case_nouns.get((match["name"], match["case"] or "nom"), nouns[match["name"]]),
        )
        if match["name"] in stats
        and match["name"] in nouns
        and (not match["case"] or match["case"] == "nom" or (match["name"], match["case"]) in case_nouns)
        else (_ for _ in ()).throw(ValueError(f"unknown document statistic phrase: {match['name']}")),
        assembled,
    )

    def replace_stat(match: re.Match[str]) -> str:
        name = match["name"]
        if name == "pages" and name not in (statistics_override or {}):
            return "[[STAT:pages]]"
        if name not in stats:
            raise ValueError(f"unknown document statistic: {name}")
        return str(stats[name])

    assembled = STAT_RE.sub(replace_stat, assembled)

    publications = completed_entries(read_bib_entries(publications_bibliography))
    groups = group_publications(publications, publication_config)
    publication_lines = []
    ordinal = 1
    for title, entries in groups:
        if not entries:
            continue
        publication_lines.extend((f"**{title}:**", ""))
        for entry in entries:
            publication_lines.append(f"[[PUBLICATION:{ordinal}]] {format_publication(entry)}")
            ordinal += 1
        publication_lines.append("")
    assembled = PUBLICATION_LIST_RE.sub("\n".join(publication_lines).strip(), assembled)

    def replace_reference(match: re.Match[str]) -> str:
        key = (match["kind"], match["id"])
        if key not in labels:
            raise ValueError(f"unknown {match['kind']} reference: {match['id']}")
        return f"[[REF:{match['kind']}:{match['id']}:{labels[key].number}]]"

    assembled = REFERENCE_RE.sub(replace_reference, assembled)

    def replace_number(match: re.Match[str]) -> str:
        key = (match["kind"], match["id"])
        if key not in labels:
            raise ValueError(f"unknown {match['kind']} reference: {match['id']}")
        return f"[[NUMBER:{match['kind']}:{match['id']}:{labels[key].number}]]"

    return NUMBER_RE.sub(replace_number, assembled).strip() + "\n"


def labels_for_kind(labels: dict[tuple[str, str], ObjectNumber], kind: str) -> list[ObjectNumber]:
    return [number for (object_kind, _), number in labels.items() if object_kind == kind]
