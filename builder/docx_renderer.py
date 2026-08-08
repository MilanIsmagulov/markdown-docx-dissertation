from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.text import WD_BREAK
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


def _number(value: str, suffix: str) -> float:
    text = str(value).strip().casefold()
    if not text.endswith(suffix):
        raise ValueError(f"expected a {suffix} value, got {value!r}")
    return float(text[: -len(suffix)])


def _set_font(style, family: str, size: str, bold: bool | None = None) -> None:
    style.font.name = family
    style.font.size = Pt(_number(size, "pt"))
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), family)
    # Pandoc's reference document leaves theme font references on built-in
    # styles (for example majorHAnsi -> Aptos Display). Word may prefer those
    # references in the style gallery even when explicit font names exist.
    for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn(f"w:{attribute}"), None)


def _set_repeatable_styles(document: Document, styles: dict) -> None:
    body = styles["body"]
    normal = document.styles["Normal"]
    _set_font(normal, body["font"]["family"], body["font"]["size"])
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = float(body["line_spacing"])
    normal.paragraph_format.first_line_indent = Mm(_number(body["first_line_indent"], "mm"))
    normal.paragraph_format.space_before = Pt(_number(body["space_before"], "pt"))
    normal.paragraph_format.space_after = Pt(_number(body["space_after"], "pt"))

    mapping = (("Heading 1", "chapter"), ("Heading 2", "section"), ("Heading 3", "subsection"))
    for word_name, yaml_name in mapping:
        spec = styles["headings"][yaml_name]
        style = document.styles[word_name]
        _set_font(style, spec["font"]["family"], spec["font"]["size"], spec["font"].get("bold"))
        style.font.color.rgb = None
        style.paragraph_format.first_line_indent = Mm(0)
        style.paragraph_format.space_before = Pt(_number(spec.get("space_before", "0pt"), "pt"))
        style.paragraph_format.space_after = Pt(_number(spec.get("space_after", "0pt"), "pt"))
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.page_break_before = bool(spec.get("page_break_before", False))
        style.paragraph_format.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER if spec.get("alignment") == "center" else WD_ALIGN_PARAGRAPH.LEFT
        )

    if "Equation" not in document.styles:
        equation = document.styles.add_style("Equation", WD_STYLE_TYPE.PARAGRAPH)
    else:
        equation = document.styles["Equation"]
    equation.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    equation.paragraph_format.first_line_indent = Mm(0)
    equation.paragraph_format.space_before = Pt(6)
    equation.paragraph_format.space_after = Pt(6)
    _set_font(equation, styles["math"]["font"], body["font"]["size"])

    for name, alignment, bold in (
        ("TOC Heading", WD_ALIGN_PARAGRAPH.CENTER, True),
        ("TOC 1", WD_ALIGN_PARAGRAPH.LEFT, False),
        ("TOC 2", WD_ALIGN_PARAGRAPH.LEFT, False),
        ("TOC 3", WD_ALIGN_PARAGRAPH.LEFT, False),
    ):
        if name not in document.styles:
            continue
        style = document.styles[name]
        _set_font(style, body["font"]["family"], body["font"]["size"], bold)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.alignment = alignment
        style.paragraph_format.first_line_indent = Mm(0)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(6 if name == "TOC Heading" else 0)


def _add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))


def _add_title_paragraph(document: Document, text: str, *, alignment, bold=False, before=0, after=0):
    paragraph = document.add_paragraph()
    paragraph.alignment = alignment
    paragraph.paragraph_format.first_line_indent = Mm(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), "Times New Roman")
    return paragraph


def _prepend_title_page(document: Document, metadata: dict) -> None:
    organization = metadata.get("organization", {}).get("full_name", "")
    author = metadata.get("author", {}).get("full_name", "")
    specialty = metadata.get("specialty", {})
    supervisor = metadata.get("supervisor", {})
    elements = []

    if organization:
        elements.append(_add_title_paragraph(document, organization, alignment=WD_ALIGN_PARAGRAPH.CENTER))
    elements.append(
        _add_title_paragraph(
            document,
            "На правах рукописи",
            alignment=WD_ALIGN_PARAGRAPH.RIGHT,
            before=18,
        )
    )
    if author:
        elements.append(
            _add_title_paragraph(document, author, alignment=WD_ALIGN_PARAGRAPH.CENTER, bold=True, before=54, after=24)
        )
    elements.append(
        _add_title_paragraph(
            document,
            str(metadata.get("title", "")).upper(),
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
            bold=True,
            after=30,
        )
    )
    elements.append(
        _add_title_paragraph(
            document,
            metadata.get("document_type", "ДИССЕРТАЦИЯ"),
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
            bold=True,
            after=12,
        )
    )
    if specialty.get("code"):
        elements.append(
            _add_title_paragraph(
                document,
                f"по специальности {specialty['code']}",
                alignment=WD_ALIGN_PARAGRAPH.CENTER,
            )
        )
    if specialty.get("name"):
        elements.append(
            _add_title_paragraph(document, specialty["name"], alignment=WD_ALIGN_PARAGRAPH.CENTER, after=12)
        )
    if metadata.get("degree"):
        elements.append(
            _add_title_paragraph(
                document,
                f"на соискание ученой степени {metadata['degree']}",
                alignment=WD_ALIGN_PARAGRAPH.CENTER,
            )
        )
    supervisor_text = " ".join(
        item for item in (supervisor.get("degree", ""), supervisor.get("title", ""), supervisor.get("full_name", "")) if item
    )
    if supervisor_text:
        elements.append(
            _add_title_paragraph(
                document,
                f"Научный руководитель:\n{supervisor_text}",
                alignment=WD_ALIGN_PARAGRAPH.RIGHT,
                before=42,
            )
        )
    place = " ".join(str(item) for item in (metadata.get("city", ""), metadata.get("year", "")) if item)
    last = _add_title_paragraph(document, place, alignment=WD_ALIGN_PARAGRAPH.CENTER, before=72)
    last.add_run().add_break(WD_BREAK.PAGE)
    elements.append(last)

    body = document._element.body
    for index, paragraph in enumerate(elements):
        body.insert(index, paragraph._p)


def _request_field_updates(document: Document) -> None:
    settings = document.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def render_docx(project_root: Path, markdown_path: Path, output_path: Path) -> Path:
    pandoc = shutil.which("pandoc")
    if pandoc is None:
        raise RuntimeError("Pandoc is required for the DOCX math backend")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        [
            pandoc,
            str(markdown_path),
            "--from=markdown+tex_math_dollars",
            "--to=docx",
            "--toc",
            "--toc-depth=3",
            "--metadata=lang:ru-RU",
            "--metadata=toc-title:ОГЛАВЛЕНИЕ",
            "--output",
            str(output_path),
        ],
        cwd=project_root,
        check=True,
    )
    styles = yaml.safe_load((project_root / "config" / "styles.yaml").read_text(encoding="utf-8-sig"))
    metadata = yaml.safe_load((project_root / "config" / "metadata.yaml").read_text(encoding="utf-8-sig")) or {}
    document = Document(output_path)
    page = styles["document"]
    for section in document.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.left_margin = Mm(_number(page["margins"]["left"], "mm"))
        section.right_margin = Mm(_number(page["margins"]["right"], "mm"))
        section.top_margin = Mm(_number(page["margins"]["top"], "mm"))
        section.bottom_margin = Mm(_number(page["margins"]["bottom"], "mm"))
        section.header_distance = Mm(_number(page["header_distance"], "mm"))
        section.footer_distance = Mm(_number(page["footer_distance"], "mm"))
        section.different_first_page_header_footer = True
        section.first_page_header.paragraphs[0].clear()
        page_numbers = styles.get("page_numbers", {})
        if page_numbers.get("position", "top-center") == "top-center":
            header = section.header
            header.paragraphs[0].clear()
            _add_page_number(header.paragraphs[0])
        else:
            footer = section.footer
            footer.paragraphs[0].clear()
            _add_page_number(footer.paragraphs[0])
    _set_repeatable_styles(document, styles)
    _prepend_title_page(document, metadata)
    _request_field_updates(document)
    for paragraph in document.paragraphs:
        if paragraph._p.xpath(".//m:oMathPara"):
            paragraph.style = document.styles["Equation"]
        if paragraph.style.name in {"Heading 1", "Heading 2", "Heading 3"}:
            for run in paragraph.runs:
                run.font.name = body_family = styles["body"]["font"]["family"]
                rpr = run._element.get_or_add_rPr()
                fonts = rpr.rFonts
                if fonts is None:
                    fonts = OxmlElement("w:rFonts")
                    rpr.insert(0, fonts)
                for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
                    fonts.set(qn(f"w:{attribute}"), body_family)
                for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                    fonts.attrib.pop(qn(f"w:{attribute}"), None)
    document.core_properties.title = "Демонстрационная сборка диссертации"
    document.core_properties.subject = "Markdown to DOCX compiler preview"
    document.save(output_path)
    return output_path
